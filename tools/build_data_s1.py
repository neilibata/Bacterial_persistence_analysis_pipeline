#!/usr/bin/env python3
"""Assemble Data S1: the tabulated data behind every figure, as one workbook.

Science asks for "tabulated data underlying the figures ... as an additional
supplementary data file (called 'data S1') in a standard machine-readable
format ... Organize this as a multi-sheet Excel document, indicating to which
figure panel each of the datasets corresponds. Include P values here as well."

Every figure already writes a `Figure_<slot>_data.pkl` carrying each panel as
a tidy frame, and `figure_source.to_workbook` already renders them as that
workbook -- but nothing called it, so the file was never produced. This is the
caller.

    python tools/build_data_s1.py                       # results/ -> Data_S1_figure_data.xlsx
    python tools/build_data_s1.py --results-root RUN --out DataS1.xlsx
    python tools/build_data_s1.py --check               # report coverage, write nothing

Slots come from tools/figure_sets.py, so a data object left behind under a
retired slot cannot creep into the submission, and a live figure with no data
object is named rather than silently skipped.
"""
from __future__ import annotations

import argparse
import pickle
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

#: Excel's hard limit, minus the header row.
XL_MAX_ROWS = 1_048_575

#: slot,panel -> caption panel letter(s) and a one-line description. The
#: journal asks the workbook to indicate "to which figure panel each of the
#: datasets corresponds", and an internal panel key such as `per_sample` does
#: not. Authored against the captions; a sheet missing from it is reported
#: rather than shipped unlabelled.
PANEL_MAP_CSV = ROOT / "tools" / "data_s1_panels.csv"

#: Significant figures kept for float columns. The run writes full double
#: precision, which is ~27% of the file in digits no reader needs; six keeps
#: every quoted statistic, including p-values down to 1e-94, intact.
SIG_FIGURES = 6

#: Columns whose name reads as a p-value but which are not one, anywhere they
#: appear. Established by range and by correlation against the sheet's own test
#: p: `p_mh` IS the Mantel-Haenszel test p (r = 0.84 against `pval`, and it
#: falls between the two arm prevalences in only 6% of rows), while these two
#: are the pooled arm prevalences.
GLOBAL_RENAMES = {"p_ctrl_mh": "prevalence_ctrl_mh",
                  "p_case_mh": "prevalence_case_mh"}

#: The same, where the name is only ambiguous in one place. F2/S19's `p` and
#: S18's `p_frac` are positions along the hunter-gatherer-to-industrial axis
#: and run outside [0, 1]; S25's `p`/`p_case` are prevalences.
PANEL_RENAMES = {
    ("F2", "axis_angle_samples"): {"p": "axis_position"},
    ("S19", "study_centroids"):   {"p": "axis_position"},
    ("S18", "bins"):              {"p_frac": "axis_position_fraction"},
    ("S25", "all_taxa"):          {"p": "prevalence_ctrl",
                                   "p_case": "prevalence_case"},
}


def _siground(col, sig: int = SIG_FIGURES):
    """Round a float column to `sig` significant figures, vectorised.

    Significant figures, not decimal places: a prevalence of 0.0045782 and a
    richness of 145.71095 both need rounding, and at different scales.
    Non-finite values and exact zeros pass through.

    The scale-and-round trick overflows on the extreme tail: for a p-value of
    1e-305 the factor is 10**310, which is `inf`, and `round(a * inf) / inf` is
    **nan** -- so the naive version silently turned the smallest p-values in
    the file into blanks. Those are rounded by exact decimal formatting
    instead, which is correct at any magnitude; there are few enough of them
    that the elementwise path costs nothing.
    """
    import numpy as np
    a = col.to_numpy(dtype=float, copy=True)
    ok = np.isfinite(a) & (a != 0)
    if not ok.any():
        return a
    mag = np.floor(np.log10(np.abs(a[ok])))
    # Keep the factor well inside the double range at both ends.
    safe = np.abs(sig - 1 - mag) <= 290
    idx = np.flatnonzero(ok)
    vec, tail = idx[safe], idx[~safe]
    if vec.size:
        f = np.power(10.0, sig - 1 - mag[safe])
        a[vec] = np.round(a[vec] * f) / f
    for i in tail:
        a[i] = float(f"%.{sig}g" % a[i])
    return a


def _prepare(df, slot: str, panel: str, sig: int = SIG_FIGURES):
    """One sheet, ready to write: renamed, annotated and rounded."""
    import numpy as np
    import pandas as pd

    out = df.copy()
    ren = dict(GLOBAL_RENAMES)
    ren.update(PANEL_RENAMES.get((slot, panel), {}))
    out = out.rename(columns={k: v for k, v in ren.items()
                              if k in out.columns})

    # A p-value that underflowed to exactly 0 reads as "not computed". Where a
    # log10 companion carries the value, say so beside it rather than leave the
    # zero bare; the true value is below double precision and cannot be
    # written as a float at all.
    for pcol in ("pval", "p", "p_raw", "p_mh"):
        if pcol not in out.columns:
            continue
        v = pd.to_numeric(out[pcol], errors="coerce")
        under = v.eq(0.0)
        if not under.any():
            continue
        log_c = next((c for c in out.columns
                      if str(c).lower() in (f"log10_{pcol}", "log10_pval",
                                            "neg_log10_pval")), None)
        note = (f"below double precision; see {log_c}" if log_c
                else "below double precision (< 1e-308)")
        out[f"{pcol}_note"] = np.where(under, note, "")

    for c in out.columns:
        if pd.api.types.is_float_dtype(out[c]):
            out[c] = _siground(out[c], sig)
    return out


#: Column names that genuinely hold a probability, once the impostors above
#: have been renamed. `p_mh` belongs here: it is the Mantel-Haenszel test p.
_P_EXACT = {"p", "q", "pval", "p_mh", "p_raw", "p_corr", "boot_p"}
_P_PREFIX = ("p_meanz", "p_energy", "p_iw", "p_binom", "p_rvc", "p_gmw",
             "p_raw_vs_gmw", "q_", "p_value", "p_perm", "p_adj")


def _is_pvalue_column(name: str) -> bool:
    n = str(name)
    return n in _P_EXACT or n.startswith(_P_PREFIX) or n.startswith("p_Obese") \
        or (n.startswith("p_") and " " in n) or n.startswith("p_CRC") \
        or n.startswith("p_IBD") or n.startswith("p_Glucose")


def pvalue_violations(objs, tol: float = 1e-9):
    """Every p/q column with a value outside [0, 1].

    A p of 2.0 was reaching the deposited clade contrast table, because a
    two-sided permutation p built as 2 x min(tail, tail) exceeds 1 whenever
    both tails do. That is the kind of thing a reader notices and nothing in
    the tree was checking, so it is checked here, against the submitted
    columns, after the prevalence-named impostors have been renamed away.
    """
    import numpy as np
    import pandas as pd
    out = []
    for obj in objs:
        slot = str(obj.get("slot") or "")
        for panel in obj["panel_order"]:
            df = _prepare(obj["panels"][panel], slot, panel)
            for c in df.columns:
                if not _is_pvalue_column(c):
                    continue
                v = pd.to_numeric(df[c], errors="coerce").to_numpy(float)
                v = v[np.isfinite(v)]
                if not v.size:
                    continue
                bad = (v < -tol) | (v > 1 + tol)
                if bad.any():
                    out.append((slot, panel, str(c), float(v.min()),
                                float(v.max()), int(bad.sum())))
    return out


def _label(slot: str) -> str:
    """The figure's journal label: `S16` -> `Fig. S16`, `F2` -> `Fig. 2`."""
    return f"Fig. {slot[1:]}" if slot.startswith("F") else f"Fig. {slot}"


def load_panel_map():
    """{(slot, panel): (caption_panel, selects_by, describes)}."""
    import pandas as pd
    if not PANEL_MAP_CSV.exists():
        print(f"  [WARN] {PANEL_MAP_CSV.name} missing; sheets will carry no "
              f"caption-panel letter")
        return {}
    m = pd.read_csv(PANEL_MAP_CSV).fillna("")
    return {(str(r.slot), str(r.panel)):
            (str(r.caption_panel), str(r.selects_panel_by), str(r.describes))
            for r in m.itertuples(index=False)}


#: Health contrasts no figure panel plots. Section 3 tests nine; Figure 5E and
#: Fig. S25 show four, and which four is decided from the data by
#: `figures/figure5.py::MIN_STRATA_TO_PLOT` -- ASCVD and Metabolic Syndrome
#: rest on a single study each and cannot be tested at all, and Fatty Liver
#: and Hypertension fall below the floor too. The panel objects carry every
#: contrast the section computed, so those rows reached the workbook anyway:
#: Data S1 tabulated, under a panel letter, contrasts that panel does not
#: show. They are dropped here, and the drop is reported rather than silent.
#:
#: Withdrawn everywhere: no panel of any figure plots these.
UNPLOTTED_CONDITIONS = ("ASCVD", "Fatty Liver", "Hypertension",
                        "Metabolic Syndrome")

#: The pooled "Unhealthy" contrast is a different case, so it is dropped per
#: slot rather than globally. It is a reported contrast, not a withdrawn one:
#: `config.HEALTH_PLOT_SKIP` holds it out of Figure 5E and Fig. S25 because
#: its cases ARE the cases of the four focal conditions -- a replicate rather
#: than a fifth condition. Fig. S26 does plot it, and keeps its rows. Dropping
#: it from the two figures that do not draw it is what makes every panel sheet
#: hold exactly what its panel shows.
SLOT_UNPLOTTED_CONDITIONS = {
    "F5":  ("Unhealthy",),
    "S25": ("Unhealthy",),
}

#: Column names that carry the health contrast across the panel frames.
_CONDITION_COLUMNS = ("Condition", "condition", "cond", "contrast")


def drop_unplotted_conditions(frame, dropped: dict, slot: str = ""):
    """Rows for contrasts the slot's panels do not plot, removed and counted.

    A frame with no condition column comes back untouched, so this is safe to
    apply to every panel of every figure rather than to a list of slots that
    would have to be kept in step with the manifest.
    """
    col = next((c for c in _CONDITION_COLUMNS if c in frame.columns), None)
    if col is None:
        return frame
    drop = set(UNPLOTTED_CONDITIONS) | set(SLOT_UNPLOTTED_CONDITIONS.get(slot, ()))
    mask = frame[col].isin(drop)
    if not int(mask.sum()):
        return frame
    for cond, k in frame.loc[mask, col].value_counts().items():
        dropped[str(cond)] = dropped.get(str(cond), 0) + int(k)
    return frame.loc[~mask]


def write_workbook(objs, out_path, extras_csv_dir=None,
                   sig: int = SIG_FIGURES, stats_path=None):
    """The submitted spreadsheet: one sheet per figure panel, plus a guide.

    Differs from `figure_source.to_workbook` in four ways the journal asks
    for, all of them post-processing over the same data objects:

    * panels held in `extras` because they are not a plain table are expanded
      into sheets (`tools/data_s1_extras.py`) -- without this, Figure 4 ships
      two sheets for an eight-panel figure;
    * every sheet is tagged with the caption panel letter(s) it underlies, and
      with the column that selects the panel where one sheet spans several;
    * a `guide` sheet carries each figure's notes, seed and reference, which
      the data objects record and the spreadsheet otherwise drops, so a
      bootstrap interval in the file can actually be regenerated;
    * floats are rounded to `sig` significant figures, and a frame too large
      for a sheet is written as a CSV beside the workbook.
    """
    import pandas as pd
    from tools import data_s1_extras

    pmap = load_panel_map()
    index, guide, unmapped, csv_written = [], [], [], []
    dropped: dict = {}
    csv_dir = Path(extras_csv_dir) if extras_csv_dir else Path(out_path).parent

    with pd.ExcelWriter(out_path) as xl:
        for obj in objs:
            slot = str(obj.get("slot") or "")
            panels = {n: obj["panels"][n] for n in obj["panel_order"]}
            add, as_csv = data_s1_extras.expand(obj)
            panels.update(add)

            for name, frame in panels.items():
                frame = drop_unplotted_conditions(frame, dropped, slot)
                if len(frame) > XL_MAX_ROWS:
                    as_csv[name] = frame
                    continue
                sheet = f"{slot}_{name}"[:31]
                prepped = _prepare(frame, slot, name, sig)
                prepped.to_excel(xl, sheet_name=sheet, index=False)
                cap, by, desc = pmap.get((slot, name), ("", "", ""))
                if not cap:
                    unmapped.append(f"{slot}/{name}")
                index.append({
                    "figure": _label(slot),
                    "caption_panel": cap or "(unmapped)",
                    "sheet": sheet,
                    "describes": desc,
                    "selects_panel_by": by,
                    "rows": len(prepped),
                    "columns": prepped.shape[1],
                    "reference": obj.get("reference", ""),
                    "panel_key": name,
                    "source_figure": obj.get("stem", ""),
                })

            for name, frame in as_csv.items():
                frame = drop_unplotted_conditions(frame, dropped, slot)
                p = csv_dir / f"Data_S1_{slot}_{name}.csv"
                _prepare(frame, slot, name, sig).to_csv(p, index=False)
                csv_written.append(p.name)
                cap, by, desc = pmap.get((slot, name), ("", "", ""))
                index.append({
                    "figure": _label(slot),
                    "caption_panel": cap or "(unmapped)",
                    "sheet": f"-> {p.name}",
                    "describes": desc or "too large for a worksheet; shipped "
                                        "as a CSV beside this workbook",
                    "selects_panel_by": by,
                    "rows": len(frame),
                    "columns": frame.shape[1],
                    "reference": obj.get("reference", ""),
                    "panel_key": name,
                    "source_figure": obj.get("stem", ""),
                })

            guide.append({
                "figure": _label(slot),
                "reference_catalogue": obj.get("reference", ""),
                "source_figure": obj.get("stem", ""),
                "written": obj.get("created", ""),
                "notes": obj.get("notes", ""),
            })

        pd.DataFrame(index).to_excel(xl, sheet_name="index", index=False)
        pd.DataFrame(guide).to_excel(xl, sheet_name="guide", index=False)
        ts = text_statistics_frame(stats_path) if stats_path else None
        if ts is not None:
            ts.to_excel(xl, sheet_name="text_statistics", index=False)

    order = _front_load_sheets(out_path)

    n_sheets = sum(1 for r in index if not r["sheet"].startswith("->"))
    print(f"  wrote {out_path}")
    print(f"    {n_sheets} panel sheets + index + guide, "
          f"{sum(r['rows'] for r in index):,} rows, "
          f"floats at {sig} significant figures")
    print(f"    opens on: {', '.join(order[:3])}")
    if ts is not None:
        print(f"    text_statistics: {len(ts)} quoted statistic(s)")
    if dropped:
        print(f"    dropped {sum(dropped.values()):,} row(s) for contrasts no "
              f"panel plots: "
              + ", ".join(f"{k} {v:,}" for k, v in sorted(dropped.items())))
    if csv_written:
        print(f"    {len(csv_written)} companion CSV(s): "
              + ", ".join(csv_written))
    if unmapped:
        # Not a warning. A sheet with no entry in the curated panel map means
        # the tree being read numbers its figures differently from the
        # manuscript -- which is exactly what a stale results/ directory looks
        # like, and it silently produced a workbook naming S5_tercile_summary
        # for what the manuscript calls S14. The map is the contract; a build
        # that cannot satisfy it must not write a submittable file.
        # Remove the file first. The check can only run once every panel has
        # been walked, which is after the workbook is on disk -- and a
        # half-trusted Data_S1 left lying beside an error message is exactly
        # the thing someone picks up later without reading the message.
        try:
            Path(out_path).unlink(missing_ok=True)
        except OSError:
            pass
        raise SystemExit(
            f"\n  ABORT: {len(unmapped)} sheet(s) have no caption-panel "
            f"entry in {PANEL_MAP_CSV.name}:\n"
            f"  (the workbook was written and then removed)\n"
            + "".join(f"      {u}\n" for u in unmapped)
            + f"\n  The slots in this results tree disagree with the "
              f"manuscript's numbering, so the workbook would carry the "
              f"wrong figure numbers.\n"
              f"  Either --results-root points at a superseded tree, or "
              f"{PANEL_MAP_CSV.name} needs a row for a genuinely new panel.\n"
              f"  Check the tree against tools/additional_figures_manifest.txt "
              f"before rebuilding.\n")
    return out_path


def _slot_key(slot: str):
    """Main-text figures first (F1..F5), then the supplement in order.

    Not every data object is a numbered slot: `make_demographic_associations`
    emits under "demographic" on purpose, so that the withdrawn S32 figure
    stays reproducible from its pickle. Sort those last rather than letting
    int() raise on them.
    """
    tail = slot[1:]
    return (0 if slot.startswith("F") else 1,
            int(tail) if tail.isdigit() else 10 ** 6, slot)


#: The reference the paper reports. Every figure is drawn once per reference,
#: so each slot has up to four data objects and the choice between them decides
#: what the submitted spreadsheet contains. Sorting by path put UHGG first.
REF_PREFERENCE = ("combined", "lev", "UHGG", "chimp")


def reference_of(path: Path, obj: dict) -> str:
    """Which reference actually drew this object.

    **The path is the authority, not `obj["reference"]`.** `figure_source.emit`
    defaults that field to "combined" and most call sites never override it, so
    very nearly every object on disk claims to be the combined one whatever
    drew it. Trusting the claim selected the UHGG copy for 16 of the 20
    contested slots while reporting them as combined -- wrong data under a
    truthful-looking label, which is the worst of the available failures.

    The claim is consulted only when the path names no reference, which is the
    case for the main-text figures under `main_text_figures/` (F1 is emitted as
    "metadata").
    """
    for r in REF_PREFERENCE:
        if r in path.parts or f"Figure3_{r}" in path.parts:
            return r
    named = str(obj.get("reference") or "")
    return named if named else "unknown"


def _pick_key(path: Path, obj: dict):
    """Rank candidate objects for one slot: the reported reference first.

    A figure built in a scratch directory loses to the same figure at the top
    of the view, which is the copy the document shows.
    """
    parts = set(path.parts)
    ref = reference_of(path, obj)
    rank = (REF_PREFERENCE.index(ref) if ref in REF_PREFERENCE
            else len(REF_PREFERENCE))
    # S4 and S23 are drawn both before and after the depth reinstatement; the
    # shipped panels are the post-reinstatement ones.
    post = 0 if "depth_audit_post_reinstatement" in parts else 1
    # main_text_figures/_figure3_build and _figure4_build are real build
    # directories; figure3.py copies the finished object up to the top of the
    # view, so the top-level copy is the one the document shows.
    scratch = 1 if any(p.startswith("_") for p in path.parts) else 0
    # Newest first among otherwise equal candidates, so running a builder
    # straight into its scratch directory cannot leave a stale top-level copy
    # winning on an alphabetical tie-break.
    try:
        age = -path.stat().st_mtime
    except OSError:
        age = 0.0
    return (rank, post, scratch, age, str(path))


#: The subtrees a finished run writes figures into. `results_pre_*` and
#: `results_stale_*` copies sit beside these and must never be read: they hold
#: data objects for the same slots from superseded runs.
SHIPPED_SUBTREES = ("results", "main_text_figures",
                    "supplementary_figures", "staged_figures")


#: Name patterns of superseded run trees. These sit beside `results/` and hold
#: data objects for the same slots from earlier runs; a stale
#: `depth_audit_post_reinstatement` object from one of them outranks a fresh
#: `depth_audit` object, so they must not be reachable even by the fallback.
SUPERSEDED_PREFIXES = ("results_pre_", "results_stale_", "results_partial_")

#: Sheets that orient a reader, in the order the workbook should open.
FRONT_SHEETS = ("index", "guide", "text_statistics")

#: Where the pipeline writes its register of reported statistics, relative to
#: a results root.
STATS_RELPATH = Path("results") / "combined" / "stats_report.csv"

#: The statistics the manuscript quotes in its running text, by tag. Curated,
#: because the pipeline records no link between a statistic and the sentence
#: that cites it -- each entry here was matched to the text by its value.
#: Deliberately short: this is the set a reader needs in order to check a
#: number they have just read, not the 1,078-row register.
TEXT_STATS_CSV = ROOT / "tools" / "text_statistics.csv"

#: Columns worth shipping. `stats_report.csv` carries about fifty, most
#: specific to one analysis and empty everywhere else.
TEXT_STATS_COLUMNS = ("tag", "quoted_as", "description", "test", "value",
                      "p_value", "ci_lo", "ci_hi", "n1", "n2", "effect_size")


def text_statistics_frame(stats_path):
    """The statistics quoted in the manuscript text, with their full values.

    Figure panels carry their own P values, so a reader checking a number on
    a figure already has it. A number quoted in the running text had nowhere
    in Data S1 to be looked up -- seven of the main text's statistics were in
    the pipeline's register and two of those only on a figure sheet.

    Attribution here is to the sentence, not to a figure: `quoted_as` says
    where in the manuscript the number appears. That is a fact about the
    text, which is why it is curated rather than derived.
    """
    import pandas as pd
    if not TEXT_STATS_CSV.exists():
        print(f"  [WARN] {TEXT_STATS_CSV.name} missing; no text_statistics "
              f"sheet")
        return None
    if not Path(stats_path).exists():
        print(f"  [WARN] {stats_path} not found; no text_statistics sheet. "
              f"Pass --stats to point at the run's stats_report.csv.")
        return None
    want = pd.read_csv(TEXT_STATS_CSV).fillna("")
    reg = pd.read_csv(stats_path, dtype=str, keep_default_na=False)
    out = want.merge(reg, on="tag", how="left", indicator=True)
    missing = out.loc[out["_merge"] == "left_only", "tag"].tolist()
    if missing:
        print(f"  [WARN] {len(missing)} quoted statistic(s) not in "
              f"{Path(stats_path).name} -- the text cites a number this run "
              f"did not produce:")
        for t in missing:
            print(f"           {t}")
    out = out.loc[out["_merge"] == "both"]
    cols = [c for c in TEXT_STATS_COLUMNS if c in out.columns]
    return out[cols].reset_index(drop=True) if len(out) else None


def _front_load_sheets(out_path, names=FRONT_SHEETS) -> list:
    """Put the orienting sheets first, and report the resulting order.

    Both pandas and openpyxl order sheets by when they were written, and the
    index cannot be written until every panel has been walked -- so it landed
    last, and the workbook opened on `F1_countries` with no map. Reordering
    the finished file is cheaper than holding every frame in memory to write
    the index first.
    """
    from openpyxl import load_workbook
    wb = load_workbook(out_path)
    front = [wb[n] for n in names if n in wb.sheetnames]
    if not front:
        return wb.sheetnames
    wb._sheets = front + [ws for ws in wb._sheets if ws not in front]
    order = list(wb.sheetnames)
    wb.save(out_path)
    return order


def _scan_roots(results_root: Path):
    """The directories to search under `results_root`.

    A run tree is scanned through its shipped subtrees only. A path that is
    itself one of those (or holds no such subtree) is scanned directly, so
    pointing this at a single reference's output still works -- but superseded
    trees are excluded from that fallback, which is otherwise exactly how they
    get back in.
    """
    roots = [results_root / s for s in SHIPPED_SUBTREES
             if (results_root / s).is_dir()]
    if roots:
        return roots
    if any(results_root.name.startswith(p) for p in SUPERSEDED_PREFIXES):
        print(f"  [WARN] {results_root.name} looks like a superseded run tree; "
              f"refusing to build Data S1 from it")
        return []
    return [results_root]


def collect(results_root: Path):
    """{slot: (object, path)} for every data object under `results_root`.

    Keeps the best candidate per slot by `_pick_key` and reports the rest, so
    a run that drew a figure against four references cannot decide the
    submission by filesystem order.
    """
    cands = {}
    paths = sorted(p for r in _scan_roots(results_root)
                   for p in r.rglob("figure_data/Figure_*_data.pkl"))
    for p in paths:
        try:
            obj = pickle.load(open(p, "rb"))
        except Exception as exc:
            print(f"  [skip] {p}: {exc}")
            continue
        slot = str(obj.get("slot") or "")
        if not slot:
            continue
        cands.setdefault(slot, []).append((p, obj))

    found, alt = {}, {}
    for slot, items in cands.items():
        items.sort(key=lambda it: _pick_key(it[0], it[1]))
        best_p, best_obj = items[0]
        found[slot] = (best_obj, best_p)
        if len(items) > 1:
            alt[slot] = [p for p, _o in items[1:]]
    return found, alt


def split_workbook(objs, full_path, limit_mb, sig, prefix="Data_S1",
                   stats_path=None):
    """Re-emit the workbook as parts, each under a portal's per-file limit.

    Science accepts 10 MB per file, and the complete workbook is well over it,
    so the submitted form is a lettered series while the complete one goes to
    the data deposit. Slots stay in figure order and no slot is split across
    parts, so a reader looking for Fig. S14 opens one file.

    Sizes come from the compressed bytes each sheet actually occupies in
    `full_path`, rather than from a row count: the sheets differ by an order
    of magnitude in bytes per row, and guessing put the main-text part over
    the limit.
    """
    import zipfile
    from openpyxl import load_workbook

    wb = load_workbook(full_path, read_only=True)
    names = wb.sheetnames
    wb.close()
    per_sheet = {}
    with zipfile.ZipFile(full_path) as z:
        for info in z.infolist():
            m = re.match(r"xl/worksheets/sheet(\d+)\.xml$", info.filename)
            if m and int(m.group(1)) - 1 < len(names):
                per_sheet[names[int(m.group(1)) - 1]] = info.compress_size

    by_slot = {}
    for sheet, nbytes in per_sheet.items():
        m = re.match(r"(F\d+|S\d+)_", sheet)
        if m:
            by_slot[m.group(1)] = by_slot.get(m.group(1), 0) + nbytes

    ordered = [o for o in objs if str(o.get("slot") or "") in by_slot]
    ordered.sort(key=lambda o: _slot_key(str(o.get("slot"))))

    limit = limit_mb * 1024 * 1024
    groups, cur, tot = [], [], 0
    for obj in ordered:
        n = by_slot[str(obj["slot"])]
        if cur and tot + n > limit:
            groups.append(cur)
            cur, tot = [], 0
        cur.append(obj)
        tot += n
    if cur:
        groups.append(cur)

    out_dir = Path(full_path).parent
    written = []
    for i, group in enumerate(groups):
        letter = chr(ord("A") + i)
        part = out_dir / f"{prefix}{letter}.xlsx"
        write_workbook(group, str(part), sig=sig, stats_path=stats_path)
        slots = [str(o["slot"]) for o in group]
        mb = part.stat().st_size / 1024 / 1024
        flag = "" if mb <= limit_mb else "   *** OVER THE LIMIT ***"
        print(f"    {part.name}: {slots[0]}-{slots[-1]}, "
              f"{len(slots)} figure(s), {mb:.2f} MB{flag}")
        written.append(part)
    return written


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--results-root", type=Path, default=ROOT,
                    help="run tree to scan (default: the repo). Only the "
                         "shipped subtrees below are read, so a superseded "
                         "results_pre_*/results_stale_* copy cannot be picked up")
    ap.add_argument("--out", type=Path, default=None)
    ap.add_argument("--check", action="store_true",
                    help="report coverage and exit non-zero if incomplete")
    ap.add_argument("--split-mb", type=float, default=None,
                    help="also write the workbook as a lettered series whose "
                         "parts each stay under this many MB, for a portal "
                         "with a per-file limit. --out stays the complete "
                         "workbook, which is what the data deposit gets")
    ap.add_argument("--split-prefix", default="Data_S1",
                    help="basename for the split parts (default: Data_S1, "
                         "giving Data_S1A.xlsx, Data_S1B.xlsx, ...)")
    ap.add_argument("--sig-figures", type=int, default=SIG_FIGURES,
                    help=f"significant figures kept for float columns "
                         f"(default {SIG_FIGURES}); the run writes full double "
                         f"precision, most of which no reader needs")
    ap.add_argument("--stats", type=Path, default=None, metavar="CSV",
                    help=f"the run's statistics register, for the "
                         f"text_statistics sheet (default: "
                         f"{STATS_RELPATH} under --results-root)")
    a = ap.parse_args()

    if not a.results_root.exists():
        print(f"no such results root: {a.results_root}", file=sys.stderr)
        return 2

    from persistence import figure_source
    from tools import figure_sets

    live = {s for s, _r, _st in figure_sets.manifest()} | {f"F{i}" for i in range(1, 6)}
    found, alt = collect(a.results_root)

    # Several references draw the same slot, so alternatives are expected
    # rather than a fault. Name the reference each sheet came from instead.
    if alt:
        n = sum(len(v) for v in alt.values())
        print(f"  {n} alternative object(s) across {len(alt)} slot(s) not used "
              f"(other references, scratch builds)")
    # Resolved from the path, never from obj["reference"] -- see reference_of.
    # The resolved value is written back onto the object so the workbook's
    # index sheet names the reference that actually drew each sheet. Leaving
    # the claim in place published "combined" against UHGG-drawn data.
    chosen = []
    for s in sorted((s for s in found if s in live), key=_slot_key):
        obj, path = found[s]
        ref = reference_of(path, obj)
        obj["reference"] = ref
        chosen.append((s, ref))
    off = [f"{s}={r}" for s, r in chosen if r not in ("combined", "metadata")]
    if off:
        print(f"  [WARN] {len(off)} slot(s) not drawn against the reported "
              f"`combined` reference: " + ", ".join(off))
        print(f"         Data S1 must not be built from these. A results tree "
              f"from a previous run is the usual cause.")

    retired = sorted(set(found) - live, key=_slot_key)
    missing = sorted(live - set(found), key=_slot_key)
    use = sorted((s for s in found if s in live), key=_slot_key)

    print(f"  {len(use)} of {len(live)} figure slots have a data object")
    if retired:
        print(f"  [WARN] {len(retired)} object(s) under retired slots, left out: "
              + ", ".join(retired))
    if missing:
        print(f"  [WARN] {len(missing)} live figure(s) with no data object: "
              + ", ".join(missing))

    oversize = [(s, n, len(found[s][0]["panels"][n]))
                for s in use for n in found[s][0]["panel_order"]
                if len(found[s][0]["panels"][n]) > XL_MAX_ROWS]
    for s, n, rows in oversize:
        print(f"  [WARN] {s}/{n} has {rows:,} rows, over Excel's sheet limit; "
              f"ship it as a CSV beside the workbook")

    invalid = pvalue_violations([found[s][0] for s in use])
    for slot, panel, col, lo, hi, n in invalid:
        print(f"  [WARN] {slot}/{panel}.{col}: {n} value(s) outside [0, 1] "
              f"(range {lo:.4g} to {hi:.4g})")
    if invalid:
        print(f"         A p-value outside [0, 1] is not a p-value. The known "
              f"cause is doubling the smaller tail of a two-sided permutation "
              f"test without clipping; health_background.py now clips, so this "
              f"clears on the next run.")

    if a.check:
        return 1 if (missing or off or oversize or invalid) else 0

    # Inside the run tree, not beside it: --results-root defaults to the repo,
    # so `.parent` put the submission workbook in the repo's parent directory.
    out = a.out or (a.results_root / "Data_S1_figure_data.xlsx")
    objs = [found[s][0] for s in use]
    stats_path = a.stats or (a.results_root / STATS_RELPATH)
    write_workbook(objs, str(out), sig=a.sig_figures,
                   stats_path=stats_path)
    print(f"  Data S1 -> {out}")
    if a.split_mb:
        print(f"  split into parts under {a.split_mb:g} MB:")
        split_workbook(objs, out, a.split_mb, a.sig_figures,
                       a.split_prefix, stats_path=stats_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
