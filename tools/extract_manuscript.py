#!/usr/bin/env python3
"""Extract a manuscript .docx into a machine-checkable JSON structure.

Emits paragraphs (with heading context), tables, figure/table captions, and an
index of every numeric claim with its surrounding sentence - the raw material
for diffing the manuscript against pipeline output.

    python tools/extract_manuscript.py --docx path/to/supplement.docx \
        --out manuscript_check/supplement.json

Numbers are captured with their sentence so a reviewer (or a diff pass) can see
what each figure refers to without reopening Word.
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import docx

# A numeric token: 1,234  0.05  1.2e-3  95%  3.4x  <0.001  ~2,000
NUM = re.compile(
    r"(?<![\w.])"
    r"(?P<op>[<>≤≥~±]\s*)?"
    r"(?P<num>\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+\.\d+(?:[eE][-+]?\d+)?|\d+)"
    r"(?P<unit>\s*(?:%|×|x\b|-fold|fold|bp|kb|Mb|Gbp|Gb|SGBs?|samples?|studies|"
    r"countries|taxa|kbp)?)"
)

CAPTION = re.compile(r"^\s*(Figure|Fig\.?|Table|Extended Data|Supplementary)\s*"
                     r"(?P<label>[A-Z]?\d+[A-Za-z]?)", re.I)


def sentences(text: str) -> list[str]:
    parts = re.split(r"(?<=[.;:!?])\s+(?=[A-Z(\[])", text)
    return [p.strip() for p in parts if p.strip()]


def extract(path: Path) -> dict:
    d = docx.Document(str(path))

    paragraphs, captions = [], []
    heading = ""
    for i, p in enumerate(d.paragraphs):
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "") if p.style is not None else ""
        if style.lower().startswith("heading") or (
                len(text) < 120 and text.isupper() and len(text.split()) > 1):
            heading = text
        rec = {"i": i, "style": style, "heading": heading, "text": text}
        paragraphs.append(rec)
        m = CAPTION.match(text)
        if m:
            captions.append({"i": i, "kind": m.group(1), "label": m.group("label"),
                             "heading": heading, "text": text})

    tables = []
    for ti, t in enumerate(d.tables):
        rows = []
        for r in t.rows:
            rows.append([c.text.strip() for c in r.cells])
        tables.append({"index": ti, "n_rows": len(rows),
                       "n_cols": max((len(r) for r in rows), default=0), "rows": rows})

    numbers = []
    for rec in paragraphs:
        for sent in sentences(rec["text"]):
            hits = list(NUM.finditer(sent))
            if not hits:
                continue
            for m in hits:
                numbers.append({
                    "para": rec["i"], "heading": rec["heading"],
                    "value": m.group("num"), "op": (m.group("op") or "").strip(),
                    "unit": (m.group("unit") or "").strip(),
                    "sentence": sent,
                })

    images = []
    with __import__("zipfile").ZipFile(path) as z:
        for n in z.namelist():
            if n.startswith("word/media/"):
                images.append({"name": Path(n).name, "bytes": z.getinfo(n).file_size})

    return {
        "source": str(path),
        "n_paragraphs": len(paragraphs), "n_tables": len(tables),
        "n_captions": len(captions), "n_numbers": len(numbers), "n_images": len(images),
        "paragraphs": paragraphs, "tables": tables, "captions": captions,
        "numbers": numbers, "images": images,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx", type=Path, required=True)
    ap.add_argument("--out", type=Path, required=True)
    args = ap.parse_args()

    data = extract(args.docx)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(data, indent=1))
    print(f"{args.docx.name}: {data['n_paragraphs']} paragraphs, {data['n_tables']} tables, "
          f"{data['n_captions']} captions, {data['n_numbers']} numeric mentions, "
          f"{data['n_images']} images -> {args.out}")


if __name__ == "__main__":
    main()
