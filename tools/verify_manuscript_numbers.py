#!/usr/bin/env python3
"""Check every cohort number in the manuscripts against the pipeline.

Absence of a stale value is not proof the right one replaced it, and a document
can be internally consistent while disagreeing with the analysis. This asserts
three things separately, and FAILS on each of them separately:

  1. no document quotes a *different* value for a quantity the pipeline
     computes -- in particular not another reference's cohort N, or a
     superseded stage's
  2. the group sizes add up
  3. the documents agree with each other about each quantity

Numbers are read from `results/<ref>/stats_report.csv` and Table S2b rather than
hard-coded, so re-running the pipeline and re-running this is enough to re-verify.
`<ref>` defaults to `config.MAIN_COHORT_REF`, the reference the manuscript
reports; it used to be hardcoded to UHGG, whose cohort (18,046) disagrees with
every N the text quotes (18,075).

    python tools/verify_manuscript_numbers.py
    python tools/verify_manuscript_numbers.py --ref UHGG
    python tools/verify_manuscript_numbers.py --docs-dir ~/Documents --suffix 081426

Exit codes: 0 all three assertions hold; 1 one of them fails; 3 an input is
absent, so nothing was verified -- which is NOT a pass.

How (1) and (3) are decided. The documents are read as prose (paragraphs) and
tables separately, because a supplementary table legitimately carries every
reference's cohort N side by side while the prose must carry only the reported
one. For each tracked quantity the pipeline supplies the correct value AND the
set of values the same quantity takes under another reference or at a
superseded stage; a document whose PROSE quotes one of those alternatives is a
conflict, named with the reference it belongs to. That set is derived from the
pipeline's own outputs, so there is nothing to keep in step by hand -- and it
is what would have caught the known failure, a 18,046 surviving in the text
after the cohort became 18,075. It does not claim to catch an arbitrary typo;
see the "scope" note it prints.
"""
from __future__ import annotations

import argparse
import re
import sys
import os as _os
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def stats(ref: str, results: Path) -> dict | None:
    """The stats report for one reference, or None if it was never written."""
    p = results / ref / "stats_report.csv"
    if not p.exists():
        return None
    d = pd.read_csv(p, dtype=str)
    t = d.columns[0]
    v = [c for c in d.columns if c.lower() in ("value", "val")][0]
    return dict(zip(d[t], d[v]))


def as_int(s) -> int | None:
    try:
        return int(float(str(s).replace(",", "")))
    except (TypeError, ValueError):
        return None


def present(text: str, value) -> bool:
    """Does this text quote this number, with or without a thousands comma?

    The trailing guard used to be ``(?![\\d,.])``, which also rejected a number
    at the end of a sentence: "4,098 non-industrial samples." did not match
    4,098, so a document that quoted the value correctly was reported as not
    quoting it at all. Reject only a continuation -- another digit, or a
    decimal point or comma followed by one.
    """
    if value is None:
        return False
    n = int(value)
    tail = r"(?!\d)(?![.,]\d)"
    pats = [rf"(?<![\d,.]){n:,}{tail}", rf"(?<![\d,.]){n}{tail}"]
    return any(re.search(p, text) for p in pats)


#: An alternative value is only treated as a possible substitute for the right
#: one when it could be mistaken for it: same number of digits, and big enough
#: not to turn up by coincidence. Without this, the chimpanzee column's 0 for
#: "of which hunter-gatherer" matched every document that contains a zero, and
#: seven coincidences buried the one real conflict.
_MIN_DISTINCTIVE = 100


def confusable(value: int, alt: int) -> bool:
    """Could `alt` be read as a stale version of `value`?"""
    return (alt != value and alt >= _MIN_DISTINCTIVE
            and len(str(int(alt))) == len(str(int(value))))


def s2b_column(ref: str) -> str:
    """The Table S2b column for a reference, from the table's own builder.

    Hardcoding "Carter (Hadza + UHGG)" is how the tool came to check the
    manuscript against a reference it does not report.
    """
    try:
        from tools.build_table_s2b_cohorts import REFS
        for r, label in REFS:
            if r == ref:
                return label
    except Exception:                                    # pragma: no cover
        pass
    return ref


def read_docs(docs_dir: Path, suffix: str) -> dict[str, tuple[str, str]]:
    """name -> (prose, table text). Kept apart on purpose: see the docstring."""
    from docx import Document
    out = {}
    for p in sorted(docs_dir.glob(f"*{suffix}.docx")):
        if p.name.startswith("~$"):
            continue
        d = Document(p)
        prose = "\n".join(x.text for x in d.paragraphs)
        tables = "\n".join(c.text for t in d.tables
                           for r in t.rows for c in r.cells)
        out[p.name] = (prose, tables)
    return out


def available_suffixes(docs_dir: Path) -> list[str]:
    """Six-digit MMDDYY suffixes with at least two documents, newest first."""
    seen: dict[str, int] = {}
    for p in docs_dir.glob("*.docx"):
        if p.name.startswith("~$"):
            continue
        m = re.search(r"_(\d{6})$", p.stem)
        if m:
            seen[m.group(1)] = seen.get(m.group(1), 0) + 1
    def _key(s):           # MMDDYY -> sortable YYMMDD
        return s[4:] + s[:4]
    return sorted((s for s, n in seen.items() if n >= 2), key=_key, reverse=True)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    # The manuscript is not in the repository; MANUSCRIPT_DOCS_DIR names
    # where it is rather than assuming one machine's layout.
    ap.add_argument("--docs-dir", type=Path,
                    default=Path(_os.environ.get("MANUSCRIPT_DOCS_DIR")
                                 or ROOT / "manuscript_docs"))
    ap.add_argument("--suffix", default=None,
                    help="document-set suffix (default: the newest set found)")
    ap.add_argument("--ref", default=None,
                    help="reference whose cohort the manuscript reports "
                         "(default: config.MAIN_COHORT_REF)")
    ap.add_argument("--results-root", type=Path, default=None,
                    help="read stats_report.csv and Table S2b from here "
                         "(e.g. a finished run's package) instead of ./results")
    args = ap.parse_args()

    results = (args.results_root.expanduser() if args.results_root
               else ROOT / "results")
    try:
        from persistence import config
        main_ref = args.ref or config.MAIN_COHORT_REF
    except Exception as exc:                             # pragma: no cover
        if not args.ref:
            print(f"  CANNOT VERIFY - config did not import ({exc}) and no "
                  "--ref was given, so the reported reference is unknown")
            return 3
        main_ref = args.ref
    col = s2b_column(main_ref)
    print(f"  reference reported by the manuscript : {main_ref}  "
          f"(Table S2b column {col!r})")

    # ── inputs, or a clear refusal ───────────────────────────────────────────
    # Each of these absences used to end the run with a bare message or a
    # KeyError. They end it with "nothing was verified" now, because a verifier
    # that cannot read its inputs has established nothing.
    absent = []
    u = stats(main_ref, results)
    if u is None:
        absent.append(f"{results / main_ref / 'stats_report.csv'}")
    s2b_path = results / "table_s2b_cohorts.csv"
    if not s2b_path.exists():
        absent.append(str(s2b_path))

    suffix = args.suffix
    avail = available_suffixes(args.docs_dir)
    if suffix is None:
        # A verifier pointed by default at a two-month-old document set checks
        # a manuscript nobody is submitting. Take the newest set on disk.
        suffix = avail[0] if avail else None
        if suffix is None:
            absent.append(f"any *_<MMDDYY>.docx set under {args.docs_dir}")
    docs = read_docs(args.docs_dir, suffix) if suffix else {}
    if suffix and not docs:
        absent.append(f"*{suffix}.docx under {args.docs_dir}")

    if absent:
        print("\n  CANNOT VERIFY - these inputs are absent, so none of the "
              "three assertions was tested:")
        for a in absent:
            print(f"    · {a}")
        print("\n  NOT A PASS. Run the pipeline (and point --results-root at "
              "its results/) first.")
        return 3

    print(f"  document set                         : {suffix}"
          + (f"   (also on disk: {', '.join(avail[1:5])})" if avail[1:] else ""))
    for n in docs:
        print(f"    {n}")

    # ── the pipeline's values ────────────────────────────────────────────────
    N = as_int(u["n_total_samples"])
    # Reported for transparency, not reconciled: these runs are excluded from
    # the cohort rather than added to it, so they enter no arithmetic identity.
    sec_runs = as_int(u.get("reinstate_disabled"))

    s2b = pd.read_csv(s2b_path)

    def s2b_row(rowlabel):
        m = s2b[s2b.iloc[:, 0].astype(str).str.contains(rowlabel, case=False,
                                                        na=False)]
        return m.iloc[0] if len(m) else None

    def s2b_val(rowlabel):
        r = s2b_row(rowlabel)
        if r is None:
            return None
        return as_int(r[col]) if col in s2b.columns else as_int(r.iloc[1])

    def s2b_alts(rowlabel) -> dict[int, str]:
        """The same row under every OTHER reference column.

        These are the values the quantity legitimately takes elsewhere, and so
        exactly the values that must not turn up in this manuscript's prose.
        """
        r = s2b_row(rowlabel)
        out: dict[int, str] = {}
        if r is None:
            return out
        for c in list(s2b.columns)[1:]:
            if c == col:
                continue
            v = as_int(r[c])
            if v is not None:
                out.setdefault(v, c)
        return out

    # "of which" anchors these to the group-composition rows; a bare
    # "industrial" also matches "Dropped: Lifestyle not Industrial/Non-Industrial"
    hg = s2b_val("of which hunter-gatherer")
    ap_ = s2b_val("of which agro-pastoralist")
    ind = s2b_val("of which industrial")

    # health split is registered post-reinstatement by the depth audit
    def health(k):
        for tag, val in u.items():
            if tag == f"depth_audit_post_Health_{k}":
                m = re.search(r"n=([\d,]+)", str(val))
                return as_int(m.group(1)) if m else None
        return None

    ni = health("Non-Industrial")
    ih, iu, ik = health("Ind-Healthy"), health("Ind-Unhealthy"), health("Ind-Unknown")

    # Alternatives for the headline N: every other reference's analytic cohort,
    # and the earlier stage of this one. A superseded cohort N surviving in the
    # prose is the failure this list exists for.
    #
    # cohort_metadata_rows and cohort_matched_to_matrix are deliberately NOT
    # here: they are different quantities the text may legitimately state, not
    # stale versions of the cohort, and treating them as conflicts flagged a
    # correct sentence about M1's row count.
    n_alts: dict[int, str] = dict(s2b_alts("Analytic cohort"))
    for other in ("UHGG", "lev", "combined", "chimp"):
        if other == main_ref:
            continue
        o = stats(other, results)
        v = as_int((o or {}).get("n_total_samples"))
        if v is not None and v != N:
            n_alts.setdefault(v, f"the {other} cohort")
    v = as_int(u.get("cohort_pre_reinstatement_samples"))
    if v is not None and v != N:
        n_alts.setdefault(v, "the pre-reinstatement cohort")

    print(f"\n  === pipeline values ===")
    facts = [
        ("analytic cohort N", N),
        ("secondary runs held back", sec_runs),
        ("hunter-gatherer", hg),
        ("agro-pastoralist", ap_),
        ("industrial", ind),
        ("non-industrial (health audit)", ni),
        ("industrial healthy", ih),
        ("industrial unhealthy", iu),
        ("industrial unknown", ik),
    ]
    for lbl, val in facts:
        print(f"    {lbl:<32}{val:>10,}" if val is not None else f"    {lbl:<32}{'?':>10}")

    bad = 0

    # ── assertion 2: the group sizes add up ─────────────────────────────────
    print(f"\n  === 2 · arithmetic ===")
    checks = [
        ("H-G + agro = non-industrial", (hg or 0) + (ap_ or 0) == ni,
         f"{hg:,} + {ap_:,} = {ni:,}" if None not in (hg, ap_, ni) else "?"),
        ("industrial health split sums", (ih or 0) + (iu or 0) + (ik or 0) == ind,
         f"{ih:,} + {iu:,} + {ik:,} = {ind:,}" if None not in (ih, iu, ik, ind) else "?"),
        ("non-industrial + industrial = N", (ni or 0) + (ind or 0) == N,
         f"{ni:,} + {ind:,} = {N:,}" if None not in (ni, ind) else "?"),
    ]
    for lbl, ok, detail in checks:
        print(f"    {'OK ' if ok else '!! '}{lbl:<34}{detail}")
        bad += (not ok)

    # ── assertions 1 and 3 ──────────────────────────────────────────────────
    watch = [
        ("analytic cohort N", N, n_alts, True),
        ("hunter-gatherer", hg, s2b_alts("of which hunter-gatherer"), False),
        ("agro-pastoralist", ap_, s2b_alts("of which agro-pastoralist"), False),
        ("industrial", ind, s2b_alts("of which industrial"), False),
        ("non-industrial", ni, {}, False),
        ("ind. healthy", ih, {}, False),
        ("ind. unhealthy", iu, {}, False),
        ("ind. unknown", ik, {}, False),
    ]

    print(f"\n  === 1 · does any document quote a DIFFERENT value? ===")
    print(f"    (prose only; a supplementary table may list every reference)")
    conflicts = []
    # quantity -> document -> the set of values that document's prose quotes
    quoted: dict[str, dict[str, set[int]]] = {}
    for lbl, val, alts, _req in watch:
        if val is None:
            print(f"    ?   {lbl:<26}the pipeline did not report this value")
            continue
        row = []
        quoted[lbl] = {}
        alts = {a: who for a, who in alts.items() if confusable(val, a)}
        for name, (prose, tables) in docs.items():
            here: set[int] = set()
            if present(prose, val):
                here.add(int(val))
            wrong = {a: who for a, who in alts.items() if present(prose, a)}
            here |= set(wrong)
            quoted[lbl][name] = here
            if wrong:
                for a, who in sorted(wrong.items()):
                    conflicts.append((lbl, name, val, a, who))
                row.append("WRONG")
            elif int(val) in here:
                row.append("prose")
            elif present(tables, val):
                row.append("table")
            else:
                row.append("  .  ")
        print(f"    {lbl:<26}{val:>9,}   " + "  ".join(f"{r:<6}" for r in row))
    print(f"    {'':<26}{'':>9}   "
          + "  ".join(f"{n.split('_')[0][:6]:<6}" for n in docs))
    if conflicts:
        print()
        for lbl, name, val, a, who in conflicts:
            print(f"    !! {name}: prose quotes {a:,} for {lbl!r}; the "
                  f"pipeline reports {val:,}  ({a:,} is {who})")
        bad += len(conflicts)
    else:
        print("\n    OK  no document's prose quotes an alternative value for "
              "any tracked quantity")

    # A quantity nothing quotes is not an error -- but the reported cohort N
    # has to appear somewhere, or the manuscript does not state its cohort.
    for lbl, val, _alts, required in watch:
        if not required or val is None:
            continue
        where = [n for n, s in quoted.get(lbl, {}).items() if int(val) in s]
        if not where:
            print(f"    !! no document's prose quotes {lbl} = {val:,}")
            bad += 1
        else:
            print(f"    OK  {lbl} = {val:,} appears in the prose of "
                  f"{len(where)}/{len(docs)} document(s)")

    print(f"\n  === 3 · do the documents agree with each other? ===")
    disagree = 0
    for lbl, val, _alts, _req in watch:
        per = {n: s for n, s in quoted.get(lbl, {}).items() if s}
        union = set().union(*per.values()) if per else set()
        if len(union) > 1:
            disagree += 1
            print(f"    !! {lbl}: " + "; ".join(
                f"{n.split('_')[0][:18]} quotes "
                + ", ".join(f"{v:,}" for v in sorted(s))
                for n, s in sorted(per.items())))
        elif len(union) == 1 and len(per) > 1:
            print(f"    OK  {lbl}: {sorted(union)[0]:,} in "
                  f"{len(per)} document(s)")
    if not disagree:
        print("    OK  no two documents quote different values for the same "
              "quantity")
    bad += disagree

    print("\n    scope: this compares the documents' prose against the "
          "pipeline's own\n    values and against the values the same "
          "quantities take under another\n    reference or at a superseded "
          "stage. It does not detect a number that\n    matches none of "
          "those -- an arbitrary typo still needs a reader.")

    print(f"\n  {'PASS - all three assertions hold' if bad == 0 else f'FAIL - {bad} problem(s)'}")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
