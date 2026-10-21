"""Render Table S12 as a Word table, ready to copy into the supplement.

The CSV has 31 columns and 54 rows, which is a data file, not a table. This
collapses it to the comparison the table exists to make -- what each method
says about the same condition-by-list cell -- by pairing the two detection
specifications inside a single cell rather than giving each its own column.

    python tools/build_table_s12_docx.py <table_s12_...csv> [out.docx]

Open the result, select the table, copy, paste into the supplement. Word
keeps the borders, shading and fonts.
"""
from __future__ import annotations

import os
import sys

import pandas as pd

# python-docx is a manuscript-tooling dependency, not a pipeline one: nothing in
# persistence imports it, it is absent from the validated plotting_env, and it
# is pinned in none of requirements.txt, requirements-pinned.txt or
# environment.yml. A hard module-level import therefore made this file
# unimportable on a clean run machine, which took tools/test_audit_fixes.py down
# with it -- the pre-run gate failed on a package the run path never loads.
#
# Importing it softly keeps the pure-pandas helpers below (_n_cell, fp, stars)
# usable anywhere, and `build()` raises a clear instruction if it is actually
# needed and missing.
try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Inches, RGBColor
    _DOCX_ERR = None
except ImportError as _exc:                      # manuscript tooling absent
    Document = WD_ORIENT = WD_TABLE_ALIGNMENT = WD_ALIGN_PARAGRAPH = None
    OxmlElement = qn = Pt = Inches = RGBColor = None
    _DOCX_ERR = _exc

FONT = "Arial"
BODY_PT = 8
HEAD_PT = 8

GREY = "F2F2F2"
INELIG = "808080"

#: Conditions in the order the figures use them.
ORDER = ["Obese", "Glucose Intolerance", "IBD", "CRC / Adenoma", "Unhealthy",
         "Fatty Liver", "Hypertension", "ASCVD", "Metabolic Syndrome"]
LISTS = ["VANISH", "BloSSUM", "VANISH+", "BloSSUM+", "VanishH+", "BlossumH+"]
DISPLAY = {"Obese": "Obesity", "VanishH+": "VANISH_H+", "BlossumH+": "BloSSUM_H+"}

CAPTION = (
 "List-level tests compared across methods and detection specifications. One row per "
 "disease condition and taxon list, grouped by list. n is the number of taxa the mean "
 "was taken over; where it is given as a fraction, the denominator is the list size and "
 "the numerator excludes taxa detected in no sample of either arm of that contrast, "
 "whose case–control shift is zero through absence of observation rather than absence "
 "of change. Mean z is the mean distance of the "
 "list from the fitted prevalence background, in units of the local spread, given under "
 "each of the two detection specifications: a 10% genome-relative breadth threshold and "
 "a flat 300,000 mapped-bases threshold. Mean z p is the two-sided bootstrap p-value "
 "under each specification. The four p HMP columns give the harmonic mean p-value across "
 "the two specifications for each of the four list-level tests \u2014 the mean z against "
 "the background curve, which is the primary test, and the importance-weighted energy "
 "distance, IW-Fisher and IW-directional permutation tests, which are supportive. "
 "Agree counts how many of the two background-referenced tests reach p < 0.05, of the "
 "number that could be computed; the energy distance is not computable where the matched "
 "background is too small. The two importance-weighted tests are one-sided in the "
 "direction that dominates the list\u2019s raw counts, given in Dom. dir., and are reported "
 "beside the others rather than counted with them.")

FOOTNOTES_HEAD = (
 "Rows in grey italic are contrasts not shown in Figure 5 or the supplementary figures: "
 "they rest on fewer than three studies with at least five cases and five controls. They "
 "are retained here so that the methods can be compared on every contrast, not only on "
 "the well-supported ones.")

#: Only written when a cell actually carries the dagger. Dropping the never-
#: detected taxa from the background (2026-10-09) removed the one disagreeing
#: cell among the reported contrasts, so on those rows the sentence would
#: otherwise read "marks the 0 cells".
FOOTNOTES_DAGGER = (
 "  \u2020 marks the {n_sign} cell{s} where the two detection specifications disagree in "
 "the sign of the mean z; {all_or_both} |mean z| below 0.12 and {is_or_are} "
 "non-significant under either specification.")

FOOTNOTES_TAIL = (
 "  The bootstrap draws 2,000 replicates, so a "
 "p-value cannot resolve below 1/2,000 and is reported as \u22640.0005.  *** p < 0.001, "
 "** p < 0.01, * p < 0.05, applied to the combined mean z.  \u2014 indicates a test that "
 "could not be computed for that cell.")


def fp(p, floor=0.0005):
    """A p-value as a journal prints it, respecting the bootstrap floor."""
    if p is None or pd.isna(p):
        return "—"
    if p <= floor:
        return f"≤{floor:g}"
    if p < 0.001:
        return f"{p:.4f}"
    if p < 0.01:
        return f"{p:.3f}"
    return f"{p:.3f}"


def stars(p):
    if p is None or pd.isna(p):
        return ""
    return "***" if p < 0.001 else "**" if p < 0.01 else "*" if p < 0.05 else ""


def _shade(cell, hex_fill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(el)


def _border(cell, edges, sz=8, color="000000"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for e in edges:
        el = OxmlElement(f"w:{e}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), color)
        borders.append(el)


def _put(cell, runs, align="left", grey=False):
    """runs: list of (text, bold, italic)."""
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT,
                   "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    p.paragraph_format.space_before = Pt(1.5)
    p.paragraph_format.space_after = Pt(1.5)
    for text, bold, ital in runs:
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(BODY_PT)
        r.bold = bold
        r.italic = ital
        if grey:
            r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)


def _n_cell(r):
    """How many taxa the cell's mean z was taken over.

    This used to print the list size alone. Since 2026-10-09 the background
    drops taxa detected in no sample of either arm (see
    config.HEALTH_BG_DROP_NEVER_DETECTED), which takes the original VANISH
    list from 124 down to 80-113 depending on the condition, so the list size
    on its own would overstate what the mean rests on. Where the two differ
    the cell reads supported/list; where they agree it reads the one number,
    which is every BloSSUM-family cell and most of the '+' cells.
    """
    n_list = r.get("n_list_10pct")
    n_sup  = r.get("n_supported_10pct")
    if pd.isna(n_list):
        return ""
    if pd.isna(n_sup) or int(n_sup) == int(n_list):
        return str(int(n_list))
    return f"{int(n_sup)}/{int(n_list)}"


COLS = [
    ("Condition", 1.02, "left"),
    ("List", 0.74, "left"),
    ("n", 0.42, "center"),
    ("Mean z\n10% / 300 kb", 0.92, "center"),
    ("Mean z p\n10% / 300 kb", 1.02, "center"),
    ("Mean z\np HMP", 0.66, "center"),
    ("Energy\np HMP", 0.64, "center"),
    ("Agree", 0.44, "center"),
    ("Dom.\ndir.", 0.40, "center"),
    ("IW-Fisher\np HMP", 0.66, "center"),
    ("IW-dir.\np HMP", 0.64, "center"),
]
#: Where the two-sided block ends and the one-sided block begins, so the
#: header can say so rather than leaving a reader to infer it.
DIR_FROM = 8


def build(csv_path, out_path=None, eligible_only=False):
    if _DOCX_ERR is not None:
        raise SystemExit(
            "python-docx is required to render Table S12 "
            f"({_DOCX_ERR}).  pip install python-docx")
    d = pd.read_csv(csv_path)
    # The pooled "Unhealthy" contrast is not reported: its cases ARE the cases
    # of the other four, so it is not an independent condition. Figure 5 and
    # S25-S31 drop it through config.HEALTH_PLOT_SKIP; this table did not, so
    # regenerating it put back a row the manuscript no longer discusses.
    _skip = set()
    try:
        sys.path.insert(0, os.path.dirname(os.path.dirname(
            os.path.abspath(__file__))))
        from persistence import config as _cfg
        _skip = set(getattr(_cfg, "HEALTH_PLOT_SKIP", ()))
    except Exception as _exc:                    # pragma: no cover
        print(f"  [warn] could not read HEALTH_PLOT_SKIP ({_exc}); "
              f"skipping no condition")
    if _skip:
        _n0 = len(d)
        d = d[~d["Condition"].isin(_skip)]
        if len(d) != _n0:
            print(f"  dropped {_n0 - len(d)} row(s) for "
                  + ", ".join(sorted(_skip)))
    d["_c"] = d["Condition"].map(lambda c: ORDER.index(c) if c in ORDER else 99)
    d["_l"] = d["Taxon_list"].map(lambda l: LISTS.index(l) if l in LISTS else 99)
    d = d.sort_values(["_l", "_c"]).reset_index(drop=True)
    if eligible_only and "eligible" in d.columns:
        d = d[d["eligible"].astype(str).str.lower() == "true"].reset_index(drop=True)

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Inches(0.6))

    def _para(text, size=8, bold_lead=None, italic=False, after=6):
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(after)
        q.paragraph_format.space_before = Pt(0)
        if bold_lead:
            r = q.add_run(bold_lead)
            r.font.name, r.font.size, r.bold = FONT, Pt(size), True
        r = q.add_run(text)
        r.font.name, r.font.size, r.italic = FONT, Pt(size), italic
        return q

    _para(CAPTION, size=8, bold_lead="Table S12. ", after=8)

    t = doc.add_table(rows=2, cols=len(COLS))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    grp = t.rows[0]
    a = grp.cells[0].merge(grp.cells[DIR_FROM - 1])
    _put(a, [("Two-sided, referenced to the background", True, False)],
         align="center")
    _shade(a, GREY)
    _border(a, ["top"], sz=12)
    b = grp.cells[DIR_FROM].merge(grp.cells[len(COLS) - 1])
    _put(b, [("One-sided, in the dominant direction", True, False)],
         align="center")
    _shade(b, GREY)
    _border(b, ["top"], sz=12)

    hdr = t.rows[1]
    for j, (name, w, al) in enumerate(COLS):
        c = hdr.cells[j]
        c.width = Inches(w)
        _put(c, [(name, True, False)], align=al)
        _shade(c, GREY)
        _border(c, ["bottom"], sz=12)

    last_list = None
    for _, r in d.iterrows():
        elig = str(r.get("eligible", "True")).lower() == "true"
        lname = DISPLAY.get(r["Taxon_list"], r["Taxon_list"])

        if lname != last_list:                      # one band per taxon list
            band = t.add_row()
            band.cells[0].merge(band.cells[len(COLS) - 1])
            _put(band.cells[0], [(lname, True, False)], align="left")
            _shade(band.cells[0], GREY)
            _border(band.cells[0], ["top"])
            last_list = lname

        row = t.add_row()
        z1, z2 = r.get("mean_z_10pct"), r.get("mean_z_300kb")
        p1, p2 = r.get("p_meanz_10pct"), r.get("p_meanz_300kb")
        hmp = r.get("p_meanz_hmp")
        agree = "" if pd.isna(r.get("n_methods_tested")) else \
            f"{int(r['n_methods_sig'])}/{int(r['n_methods_tested'])}"
        sign_ok = bool(r.get("sign_agree", True))

        cells = [
            [(DISPLAY.get(r["Condition"], r["Condition"]), False, not elig)],
            [(lname, False, not elig)],
            [(_n_cell(r), False, False)],
            [(f"{z1:+.2f} / {z2:+.2f}" if pd.notna(z1) and pd.notna(z2)
              else "—", False, False)] +
            ([("†", False, False)] if not sign_ok else []),
            [(f"{fp(p1)} / {fp(p2)}", False, False)],
            [(fp(hmp), bool(pd.notna(hmp) and hmp < 0.05), False),
             (" " + stars(hmp), bool(pd.notna(hmp) and hmp < 0.05), False)],
            [(fp(r.get("p_energy_hmp")), False, False)],
            [(agree, False, False)],
            [(str(r.get("direction", "") or "").replace("nan", ""), False, True)],
            [(fp(r.get("p_iw_fisher_hmp")), False, False)],
            [(fp(r.get("p_iw_directional_hmp")), False, False)],
        ]
        for j, runs in enumerate(cells):
            c = row.cells[j]
            c.width = Inches(COLS[j][1])
            _put(c, [x for x in runs if x[0]], align=COLS[j][2], grey=not elig)

    for j in range(len(COLS)):                      # closing rule
        _border(t.rows[-1].cells[j], ["bottom"], sz=12)

    n_sign = int((~d.get("sign_agree", pd.Series([True] * len(d)))
                  .astype(bool)).sum())
    _foot = FOOTNOTES_HEAD
    if n_sign:
        _foot += FOOTNOTES_DAGGER.format(
            n_sign=n_sign, s="" if n_sign == 1 else "s",
            all_or_both="it has" if n_sign == 1 else "all have",
            is_or_are="is" if n_sign == 1 else "are")
    _foot += FOOTNOTES_TAIL
    _para(_foot, size=7, after=0)

    out_path = out_path or os.path.splitext(csv_path)[0] + "_word.docx"
    doc.save(out_path)
    print(f"wrote {out_path}  ({len(d)} rows, {len(COLS)} columns, landscape)")
    return out_path


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    build(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
