#!/usr/bin/env python3
"""
Build Section 2 supplementary pooling tables from existing pipeline outputs.

The script summarizes the per-threshold CSVs produced by
persistence.section2_pooling into the three manuscript-facing tables shown
in the supplement:

  1. Detection-burden summary across thresholds.
  2. Pooled missing taxa for original VANISH / BloSSUM lists.
  3. Rare/extinction-candidate taxa for VANISH+ / BloSSUM+ lists.

By default this now targets the combined reference.
"""

from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import pandas as pd


DEFAULT_THRESHOLDS = ("10pct", "30pct", "50pct", "300kb", "500kb")

THRESHOLD_LABELS = {
    "10pct": "10% breadth",
    "30pct": "30% breadth",
    "50pct": "50% breadth",
    "300kb": "300 kb",
    "500kb": "500 kb",
}

POOLING_FILES = {
    "all_ind": "all_sgb_pooling_industrial_{tag}.csv",
    "all_ni": "all_sgb_pooling_ni_{tag}.csv",
    "vanish": "vanish_pooling_industrial_{tag}.csv",
    "vanish_plus": "vanish_plus_pooling_industrial_{tag}.csv",
    "blossum": "blossum_pooling_ni_{tag}.csv",
    "blossum_plus": "blossum_plus_pooling_ni_{tag}.csv",
}


def _read_csv(path: Path, required: bool = True) -> pd.DataFrame:
    if path.exists():
        return pd.read_csv(path)
    if required:
        raise FileNotFoundError(path)
    return pd.DataFrame()


def _median_n(df: pd.DataFrame) -> float:
    if df.empty or "n_needed" not in df:
        return np.nan
    return float(df["n_needed"].median())


def _det_string(macro: pd.DataFrame, list_name: str, context: str) -> str:
    sub = macro[(macro["list"] == list_name) & (macro["context"] == context)]
    if sub.empty:
        return ""
    return f"{int((sub['n_detected'] > 0).sum())}/{len(sub)}"


def _min_hard_det(macro: pd.DataFrame, list_name: str, context: str) -> float:
    sub = macro[(macro["list"] == list_name) & (macro["context"] == context)]
    if sub.empty:
        return np.nan
    return int(sub["n_detected"].min())


def _last_rank(taxonomy: object) -> str:
    if pd.isna(taxonomy):
        return ""
    parts = [p.strip() for p in str(taxonomy).split(";") if p.strip()]
    for part in reversed(parts):
        if not part.endswith("__"):
            return part
    return ""


def _taxon_label(row: pd.Series) -> str:
    taxon = str(row["taxon"])
    rank = _last_rank(row.get("taxonomy", ""))
    return f"{taxon} ({rank})" if rank else taxon


def _missing_taxa(macro: pd.DataFrame, list_name: str, context: str) -> str:
    sub = macro[(macro["list"] == list_name) & (macro["context"] == context)]
    missing = sub[sub["n_detected"] <= 0]
    if missing.empty:
        return "none"
    return "; ".join(_taxon_label(row) for _, row in missing.iterrows())


def _rare_taxa(candidates: pd.DataFrame, direction: str) -> str:
    if candidates.empty:
        return "none"
    sub = candidates[candidates["direction"] == direction]
    if sub.empty:
        return "none"
    return "; ".join(_taxon_label(row) for _, row in sub.iterrows())


def _candidate_count(candidates: pd.DataFrame, direction: str) -> int:
    if candidates.empty:
        return 0
    return int((candidates["direction"] == direction).sum())


def build_tables(ref: str, results_dir: Path, thresholds: tuple[str, ...],
                 original_ref: str | None = None) -> dict[str, pd.DataFrame]:
    """Build Tables S6-S8.

    The "original VANISH / original BloSSUM" columns describe the two published
    curated lists. Those lists were *defined* against UHGG, but every other
    column here -- the typical-SGB baselines they are compared against, and the
    VANISH+/BloSSUM+ columns -- is computed on `ref`. Evaluating the original
    lists on a different catalogue than their own baseline makes the fold-change
    comparison meaningless, so `original_ref` defaults to `ref`. Pass
    `--original-ref UHGG` to recover the old split-basis behaviour.
    """
    original_ref = original_ref or ref
    combined_main = results_dir / ref / "section2_differential" / "main"
    original_main = results_dir / original_ref / "section2_differential" / "main"
    if not combined_main.exists():
        raise FileNotFoundError(combined_main)
    if not original_main.exists():
        raise FileNotFoundError(original_main)

    summary_rows = []
    missing_rows = []
    rare_rows = []

    for tag in thresholds:
        combined_pools = {
            name: _read_csv(combined_main / template.format(tag=tag))
            for name, template in POOLING_FILES.items()
        }
        original_pools = {
            name: _read_csv(original_main / template.format(tag=tag))
            for name, template in POOLING_FILES.items()
            if name in {"vanish", "blossum"}
        }
        combined_macro = _read_csv(combined_main / f"macro_pooling_detection_{tag}.csv")
        original_macro = _read_csv(original_main / f"macro_pooling_detection_{tag}.csv")
        candidates = _read_csv(
            combined_main / f"extinction_candidates_vanish_blossum_plus_{tag}.csv",
            required=False,
        )

        summary_rows.append({
            "Threshold": THRESHOLD_LABELS.get(tag, tag),
            "Median N to detect SGB - Ind": _median_n(combined_pools["all_ind"]),
            "Median N to detect SGB - NI": _median_n(combined_pools["all_ni"]),
            "original VANISH - det.": _det_string(original_macro, "VANISH", "Industrial"),
            "original VANISH - Med N": _median_n(original_pools["vanish"]),
            "original VANISH - Min hard det.": _min_hard_det(original_macro, "VANISH", "Industrial"),
            "VANISH+ - det.": _det_string(combined_macro, "VANISH+", "Industrial"),
            "VANISH+ - Med N": _median_n(combined_pools["vanish_plus"]),
            "VANISH+ - Min hard det.": _min_hard_det(combined_macro, "VANISH+", "Industrial"),
            "original BloSSUM - det.": _det_string(original_macro, "BloSSUM", "Non-Industrial"),
            "original BloSSUM - Med N": _median_n(original_pools["blossum"]),
            "original BloSSUM - Min hard det.": _min_hard_det(original_macro, "BloSSUM", "Non-Industrial"),
            "BloSSUM+ - pooled det.": _det_string(combined_macro, "BloSSUM+", "Non-Industrial"),
            "BloSSUM+ - Med N": _median_n(combined_pools["blossum_plus"]),
            "BloSSUM+ - Min hard det.": _min_hard_det(combined_macro, "BloSSUM+", "Non-Industrial"),
        })

        missing_rows.append({
            "Threshold": THRESHOLD_LABELS.get(tag, tag),
            f"Original VANISH in Industrial (pooled; {original_ref}) - det.": _det_string(original_macro, "VANISH", "Industrial"),
            f"Original VANISH in Industrial (pooled; {original_ref}) - Pooled-missing taxa": _missing_taxa(
                original_macro, "VANISH", "Industrial"),
            f"Original BloSSUM in Non-Industrial (pooled; {original_ref}) - det.": _det_string(original_macro, "BloSSUM", "Non-Industrial"),
            f"Original BloSSUM in Non-Industrial (pooled; {original_ref}) - Pooled-missing taxa": _missing_taxa(
                original_macro, "BloSSUM", "Non-Industrial"),
        })

        rare_rows.append({
            "Threshold": THRESHOLD_LABELS.get(tag, tag),
            "VANISH+ in Industrial (pooled; combined) - det.": _det_string(combined_macro, "VANISH+", "Industrial"),
            "VANISH+ in Industrial (pooled; combined) - N candidates (<=3 hard det.)": _candidate_count(candidates, "VANISH+"),
            "VANISH+ in Industrial (pooled; combined) - Rare taxa": _rare_taxa(candidates, "VANISH+"),
            "BloSSUM+ in Non-Industrial (pooled; combined) - det.": _det_string(combined_macro, "BloSSUM+", "Non-Industrial"),
            "BloSSUM+ in Non-Industrial (pooled; combined) - N candidates (<=3 hard det.)": _candidate_count(candidates, "BloSSUM+"),
            "BloSSUM+ in Non-Industrial (pooled; combined) - Rare taxa": _rare_taxa(candidates, "BloSSUM+"),
        })

    return {
        "detection_summary": pd.DataFrame(summary_rows),
        "pooled_missing_taxa": pd.DataFrame(missing_rows),
        "rare_candidate_taxa": pd.DataFrame(rare_rows),
    }

def _write_xlsx(tables: dict[str, pd.DataFrame], out_xlsx: Path) -> None:
    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
        for sheet_name, df in tables.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)

    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    workbook = load_workbook(out_xlsx)
    header_fill = PatternFill("solid", fgColor="D9EAD3")
    header_font = Font(bold=True)
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for sheet_name, df in tables.items():
        worksheet = workbook[sheet_name]
        worksheet.freeze_panes = "B2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col_idx, col in enumerate(df.columns, start=1):
            width = min(max(len(str(col)) + 2, 12), 48)
            if not df.empty:
                width = min(max(width, int(df[col].astype(str).str.len().quantile(0.90)) + 2), 70)
            worksheet.column_dimensions[worksheet.cell(row=1, column=col_idx).column_letter].width = width
    workbook.save(out_xlsx)


def write_tables(tables: dict[str, pd.DataFrame], out_dir: Path, prefix: str) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for name, df in tables.items():
        path = out_dir / f"{prefix}_{name}.csv"
        df.to_csv(path, index=False)
        paths.append(path)
    xlsx_path = out_dir / f"{prefix}_section2_supplementary_tables.xlsx"
    _write_xlsx(tables, xlsx_path)
    paths.append(xlsx_path)
    return paths


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--ref", default="combined", help="Reference name under results/ (default: combined).")
    parser.add_argument("--results-dir", default="results", type=Path, help="Top-level results directory.")
    parser.add_argument("--out-dir", default=None, type=Path, help="Output directory. Defaults to supplement dir for ref.")
    parser.add_argument("--thresholds", nargs="+", default=list(DEFAULT_THRESHOLDS), help="Threshold tags to summarize.")
    parser.add_argument("--original-ref", default=None,
                        help="Reference for the 'original VANISH/BloSSUM' columns "
                             "(default: same as --ref, so the whole table shares one basis).")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    thresholds = tuple(args.thresholds)
    out_dir = args.out_dir or args.results_dir / args.ref / "section2_differential" / "supplement"
    tables = build_tables(args.ref, args.results_dir, thresholds, args.original_ref)
    paths = write_tables(tables, out_dir, args.ref)
    print(f"Wrote {len(paths)} supplementary table files:")
    for path in paths:
        print(f"  {path}")


if __name__ == "__main__":
    main()
